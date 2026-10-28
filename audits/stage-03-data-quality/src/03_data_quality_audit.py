"""Stage 03: read-only data-quality checks on the master workbook.

Run from project root: python audits/stage-03-data-quality/src/03_data_quality_audit.py
No cleaning or source-file modification is performed.
"""
from pathlib import Path
from collections import defaultdict
import math
import pandas as pd
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[3]
STAGE_DIR = Path(__file__).resolve().parents[1]
INPUT = REPO_ROOT / "data" / "raw" / "master excel.xlsx"
REPORT_DIR = STAGE_DIR / "reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)
SHEET_NAME = "master dataset"
KEY_COLS = ["Country Name", "Country Code", "Year"]
# These are semantic expectations, not unquestionable truths. They produce flags to review.
SHARE_COLS = ["fossil_share_elec", "fossil_share_energy", "low_carbon_share_elec",
              "low_carbon_share_energy", "Share of global annual CO₂ emissions"]
NONNEGATIVE_COLS = ["GDP", "population", "electricity_demand", "electricity_generation",
                    "energy_per_capita", "energy_per_gdp", "greenhouse_gas_emissions",
                    "Annual CO2 emissions", "CO2 emissions per capita"]
NUMERIC_EXPECTED = ["Year", "GDP", "Inflation", "Trade", "population", "carbon_intensity_elec",
                    "electricity_demand", "electricity_generation", "energy_per_capita", "energy_per_gdp",
                    "fossil_share_elec", "fossil_share_energy", "greenhouse_gas_emissions",
                    "low_carbon_share_elec", "low_carbon_share_energy", "Annual CO2 emissions",
                    "CO2 emissions per capita", "Share of global annual CO₂ emissions", "GDP Growth (%)",
                    "CO2 Growth (%)", "Transition Risk Score"]


def add_issue(rows, severity, check, column, count, detail, example=""):
    rows.append({"severity": severity, "check": check, "column": column,
                 "count": int(count), "detail": detail, "example_rows_or_values": str(example)[:500]})


def main():
    if not INPUT.exists():
        raise FileNotFoundError(f"Workbook not found: {INPUT}")
    df = pd.read_excel(INPUT, sheet_name=SHEET_NAME)
    issues = []

    # 1) Required columns and key checks
    for c in KEY_COLS:
        if c not in df.columns:
            add_issue(issues, "ERROR", "required_column_missing", c, 1, "Required key column is absent.")
        else:
            n = int(df[c].isna().sum())
            if n:
                add_issue(issues, "ERROR", "missing_key_values", c, n, "Rows cannot be fully identified without this key value.", df.index[df[c].isna()].tolist()[:10])
    if all(c in df.columns for c in KEY_COLS):
        dup = df.duplicated(["Country Code", "Year"], keep=False)
        if dup.any():
            add_issue(issues, "ERROR", "duplicate_entity_year", "Country Code + Year", int(dup.sum()),
                      "Multiple rows share a country/entity code and year; inspect before analysis.",
                      df.loc[dup, ["Country Name", "Country Code", "Year"]].head(10).to_dict("records"))
        name_per_code = df.groupby("Country Code", dropna=False)["Country Name"].nunique(dropna=True)
        bad_codes = name_per_code[name_per_code > 1]
        if len(bad_codes):
            add_issue(issues, "REVIEW", "code_maps_to_multiple_names", "Country Code", len(bad_codes),
                      "A code maps to more than one name. Could be a naming/version issue; inspect examples.", bad_codes.head(10).to_dict())
        code_per_name = df.groupby("Country Name", dropna=False)["Country Code"].nunique(dropna=True)
        bad_names = code_per_name[code_per_name > 1]
        if len(bad_names):
            add_issue(issues, "REVIEW", "name_maps_to_multiple_codes", "Country Name", len(bad_names),
                      "A name maps to more than one code. Could reflect source naming differences; inspect examples.", bad_names.head(10).to_dict())

    # 2) Numeric checks using pandas-parsed values
    for col in NUMERIC_EXPECTED:
        if col not in df.columns:
            add_issue(issues, "ERROR", "expected_numeric_column_missing", col, 1, "Expected numeric variable is absent.")
            continue
        if not pd.api.types.is_numeric_dtype(df[col]):
            nonempty = df[col].notna() & df[col].astype(str).str.strip().ne("")
            add_issue(issues, "REVIEW", "column_not_numeric_dtype", col, int(nonempty.sum()),
                      f"Pandas inferred dtype {df[col].dtype}; inspect non-empty values.", df.loc[nonempty, col].head(10).tolist())
        # Negative values for measures expected to be non-negative; not a blanket rule for growth/rates.
        if col in NONNEGATIVE_COLS and pd.api.types.is_numeric_dtype(df[col]):
            mask = df[col].notna() & (df[col] < 0)
            if mask.any():
                add_issue(issues, "REVIEW", "negative_value_in_nonnegative_measure", col, int(mask.sum()),
                          "Negative values may be invalid for this measure; verify definition and units.", df.loc[mask, ["Country Name", "Year", col]].head(10).to_dict("records"))
        if col in SHARE_COLS and pd.api.types.is_numeric_dtype(df[col]):
            mask = df[col].notna() & ((df[col] < 0) | (df[col] > 100))
            if mask.any():
                add_issue(issues, "REVIEW", "share_outside_0_100", col, int(mask.sum()),
                          "Share appears outside 0–100. Verify whether this column is stored as a fraction, percent, or different metric.", df.loc[mask, ["Country Name", "Year", col]].head(10).to_dict("records"))

    if "Year" in df.columns:
        y = pd.to_numeric(df["Year"], errors="coerce")
        mask = y.notna() & ((y % 1 != 0) | (y < 1800) | (y > 2100))
        if mask.any():
            add_issue(issues, "REVIEW", "unusual_year", "Year", int(mask.sum()),
                      "Year is non-integer or outside a broad plausible range; verify source.", df.loc[mask, ["Country Name", "Year"]].head(10).to_dict("records"))

    # 3) Inspect raw Excel cell contents: catches non-empty text inside expected numeric columns
    # that pandas may have coerced to NaN. Excel row numbers are 1-based, including header row.
    wb = load_workbook(INPUT, read_only=True, data_only=False)
    ws = wb[SHEET_NAME]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    raw_text_rows = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        for col_idx, (col, value) in enumerate(zip(headers, row), start=1):
            if col in NUMERIC_EXPECTED and isinstance(value, str) and value.strip():
                # Formula cells are noted separately; a formula string begins with '='.
                kind = "formula" if value.startswith("=") else "text_in_numeric_column"
                raw_text_rows.append({"excel_row": row_num, "column": col, "raw_value": value[:200], "cell_type": kind})
    wb.close()
    pd.DataFrame(raw_text_rows, columns=["excel_row", "column", "raw_value", "cell_type"]).to_csv(REPORT_DIR / "03_nonnumeric_cells.csv", index=False)
    if raw_text_rows:
        add_issue(issues, "REVIEW", "raw_text_in_expected_numeric_columns", "Multiple", len(raw_text_rows),
                  "Non-empty text/formula strings were found in columns expected to be numeric. Check the CSV for exact cells.", raw_text_rows[:8])

    # 4) Unexpected category values and whitespace/case variants
    cat_cols = [c for c in ["Country Name", "Country Code", "Region", "IncomeGroup", "Decoupling Status"] if c in df.columns]
    cat_rows = []
    for col in cat_cols:
        s = df[col].dropna().astype(str)
        for value, count in s.value_counts(dropna=False).items():
            stripped = value.strip()
            cat_rows.append({"column": col, "value": value, "count": int(count),
                             "has_leading_or_trailing_whitespace": value != stripped,
                             "lowercase_normalized_value": stripped.casefold()})
        whitespace_mask = df[col].notna() & (df[col].astype(str) != df[col].astype(str).str.strip())
        if whitespace_mask.any():
            add_issue(issues, "REVIEW", "category_whitespace", col, int(whitespace_mask.sum()),
                      "Leading/trailing whitespace can create false category mismatches.", df.loc[whitespace_mask, col].head(10).tolist())
    pd.DataFrame(cat_rows).to_csv(REPORT_DIR / "03_category_values.csv", index=False)

    # 5) Numeric outlier screening via IQR. Screening only; never labels values as errors automatically.
    outlier_rows = []
    for col in df.columns:
        if pd.api.types.is_numeric_dtype(df[col]) and df[col].notna().sum() >= 8:
            s = df[col].dropna()
            q1, q3 = s.quantile([0.25, 0.75])
            iqr = q3 - q1
            if pd.isna(iqr) or iqr == 0:
                continue
            low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
            mask = df[col].notna() & ((df[col] < low) | (df[col] > high))
            for idx in df.index[mask][:500]:
                outlier_rows.append({"column": col, "dataframe_row_index": int(idx),
                    "country_name": df.at[idx, "Country Name"] if "Country Name" in df.columns else "",
                    "year": df.at[idx, "Year"] if "Year" in df.columns else "",
                    "value": df.at[idx, col], "iqr_lower_fence": low, "iqr_upper_fence": high,
                    "interpretation": "Review candidate only; extreme values can be real."})
            if mask.any():
                add_issue(issues, "INFO", "iqr_outlier_candidates", col, int(mask.sum()),
                          "Pooled IQR screen across countries/years; scale differences can flag many valid observations. Not proof of error. See 03_iqr_outlier_candidates.csv.", df.loc[mask, ["Country Name", "Year", col]].head(5).to_dict("records"))
    pd.DataFrame(outlier_rows, columns=["column", "dataframe_row_index", "country_name", "year", "value", "iqr_lower_fence", "iqr_upper_fence", "interpretation"]).to_csv(REPORT_DIR / "03_iqr_outlier_candidates.csv", index=False)

    # 6) Report summary and issue log
    issue_df = pd.DataFrame(issues, columns=["severity", "check", "column", "count", "detail", "example_rows_or_values"])
    issue_df.to_csv(REPORT_DIR / "03_issue_log.csv", index=False)
    severity_counts = issue_df["severity"].value_counts().to_dict() if not issue_df.empty else {}
    lines = [
        "STAGE 03 — DATA QUALITY AUDIT (READ-ONLY)", "=" * 52,
        f"Workbook: {INPUT.name}", f"Sheet: {SHEET_NAME}", f"Shape: {len(df):,} rows x {len(df.columns)} columns", "",
        "IMPORTANT: This script does not change or clean the source workbook.",
        "REVIEW/INFO flags are prompts for investigation, not automatic proof of error.",
        "Share ranges assume 0–100 for screening only; verify units first.", "",
        f"Issue flags: {len(issue_df)}", f"Severity counts: {severity_counts}", "",
        "Files written:", "- reports/03_issue_log.csv: all checks and examples",
        "- reports/03_nonnumeric_cells.csv: raw non-empty text/formula cells in expected numeric columns",
        "- reports/03_category_values.csv: observed category values and frequencies",
        "- reports/03_iqr_outlier_candidates.csv: statistical outlier candidates, not errors", "",
        "INTERPRETATION GUIDE:",
        "1. Start with ERROR rows in the issue log; then review REVIEW flags.",
        "2. INFO IQR candidates are statistical extremes and can be legitimate.",
        "3. Pooled IQR screening across countries/years is deliberately broad and can flag valid large countries/values; use within-country or log-scale checks later if needed.",
        "4. Do not replace missing values with zero or delete flagged rows automatically.",
        "5. Source files/metadata are needed to prove merge loss or confirm units and definitions.",
    ]
    (REPORT_DIR / "03_data_quality_summary.txt").write_text("\n".join(lines), encoding="utf-8")
    print("Stage 03 completed. Read reports/03_data_quality_summary.txt first.")
    print(f"Issue log: {REPORT_DIR / '03_issue_log.csv'}")

if __name__ == "__main__":
    main()
